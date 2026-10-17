"""Build backend/data/sites.json from the EIA-860M monthly generator inventory.

    cd backend && python -m scripts.load_sites            # download latest, rebuild
    cd backend && python -m scripts.load_sites --file data/raw/august_generator2026.xlsx

Steps: download the newest EIA-860M workbook -> read the "Operating" sheet ->
keep Florida generators -> group into one row per plant -> pick 55 plants
(largest first, with regional minimums, no co-located duplicates) -> add IANA timezone -> write JSON.
Also writes every Florida plant that burns fuel or makes steam to fl_thermal_plants.json: NASA FIRMS
detections at those plants are plant heat, not wildfire (hazards/sources/firms.py).

The raw workbook goes to backend/data/raw/ (gitignored). Only the two JSON files are committed,
so the app never downloads anything at runtime. Needs: pip install -r scripts/requirements.txt
"""

import argparse
import json
import math
import re
import urllib.request
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

BACKEND_DIR = Path(__file__).resolve().parent.parent
RAW_DIR = BACKEND_DIR / "data" / "raw"
OUT_FILE = BACKEND_DIR / "data" / "sites.json"
THERMAL_FILE = BACKEND_DIR / "data" / "fl_thermal_plants.json"
# Technologies with no stack or flare: a satellite hot spot there may be a real fire.
NOT_THERMAL = {"Solar Photovoltaic", "Batteries", "Conventional Hydroelectric", "Onshore Wind Turbine", "Flywheels"}

EIA_BASE = "https://www.eia.gov"
EIA_PAGE = f"{EIA_BASE}/electricity/data/eia860m/"
USER_AGENT = "Mozilla/5.0 (PlanIT Control site loader)"
MONTHS = ["january", "february", "march", "april", "may", "june", "july",
          "august", "september", "october", "november", "december"]

TARGET_COUNT = 55

# Regions used to spread the 55 sites across the state. Panhandle is by county.
# "everglades" is inland South Florida (west of the coastal cities, north of the Keys:
# Everglades, Big Cypress, the sugar lands south of Lake Okeechobee). The rest are latitude bands.
PANHANDLE_COUNTIES = {
    "Escambia", "Santa Rosa", "Okaloosa", "Walton", "Holmes", "Washington", "Bay", "Jackson",
    "Calhoun", "Gulf", "Liberty", "Franklin", "Gadsden", "Leon", "Wakulla", "Jefferson",
}
EVERGLADES_BOX = {"lat": (25.2, 26.9), "lon": (-81.5, -80.4)}
REGION_MINIMUMS = {"panhandle": 5, "everglades": 5, "north": 5, "central": 5, "south": 4}

# Plants closer than this to a larger selected plant are the same job site (e.g. Dania Beach 7
# sits on the old Lauderdale plant), so they are skipped rather than shown as a second pin.
CO_LOCATED_MI = 1.0

# EIA spells some counties inconsistently across rows.
COUNTY_FIXES = {"Desoto": "DeSoto", "Miami Dade": "Miami-Dade", "St Johns": "St. Johns",
                "St Lucie": "St. Lucie"}

# Florida's Central time zone, as a (lon, lat) polygon. The eastern edge follows the
# Chattahoochee/Apalachicola River south from the Georgia line, then the Intracoastal
# Waterway west across Gulf County (north of it is Central, Port St. Joe and Cape San Blas
# are Eastern), then the Gulf/Bay county line to the coast. Approximate, but accurate to a
# few miles, which is enough for plant locations. Cross-checked against CENTRAL_COUNTIES below.
CENTRAL_TIME_POLYGON = [
    (-87.70, 31.10), (-85.00, 31.10), (-84.86, 30.72),            # AL/GA line to Chattahoochee
    (-84.93, 30.62), (-84.96, 30.55), (-84.99, 30.45),            # Apalachicola River
    (-85.03, 30.35), (-85.08, 30.22), (-85.12, 30.10),
    (-85.07, 29.95), (-85.05, 29.82),
    (-85.17, 29.83), (-85.22, 29.88), (-85.30, 29.97),            # ICW through Gulf County
    (-85.38, 30.03), (-85.40, 29.94),                             # Gulf/Bay county line
    (-85.60, 29.60), (-87.70, 29.60),                             # offshore, back west
]
# Counties entirely on Central time. Gulf is split, so the polygon decides it.
CENTRAL_COUNTIES = {"Escambia", "Santa Rosa", "Okaloosa", "Walton", "Holmes", "Washington",
                    "Bay", "Jackson", "Calhoun"}


# ---------- 1. Download ----------

def latest_workbook_urls() -> list[str]:
    """Generator-inventory .xlsx links on the EIA-860M page, newest first.

    The page keeps placeholder links for future months inside an HTML comment, so
    comments are stripped before searching.
    """
    req = urllib.request.Request(EIA_PAGE, headers={"User-Agent": USER_AGENT})
    html = urllib.request.urlopen(req, timeout=60).read().decode("utf-8", "replace")
    html = re.sub(r"<!--.*?-->", "", html, flags=re.DOTALL)
    links = set(re.findall(r'href="([^"]*/([a-z]+)_generator(\d{4})\.xlsx)"', html))
    ranked = sorted(links, key=lambda m: (int(m[2]), MONTHS.index(m[1])), reverse=True)
    return [EIA_BASE + path if path.startswith("/") else path for path, _, _ in ranked]


def download_latest() -> Path:
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    for url in latest_workbook_urls():
        dest = RAW_DIR / url.rsplit("/", 1)[1]
        if dest.exists():
            print(f"Using cached {dest.relative_to(BACKEND_DIR)}")
            return dest
        try:
            req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
            data = urllib.request.urlopen(req, timeout=300).read()
        except OSError as err:
            print(f"Skipping {url}: {err}")
            continue
        dest.write_bytes(data)
        print(f"Downloaded {url} ({len(data) / 1e6:.1f} MB)")
        return dest
    raise SystemExit("No EIA-860M generator workbook could be downloaded")


# ---------- 2. Read Florida generators ----------

def read_fl_generators(path: Path) -> list[dict]:
    """Rows of the Operating sheet where Plant State == FL, as dicts keyed by column name.

    The sheet starts with a title row and a blank row; the real header is the first row
    containing "Plant State". The last rows are notes, which have no numeric Plant ID.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = wb["Operating"].iter_rows(values_only=True)
    header = next(r for r in rows if r and "Plant State" in r)
    generators = []
    for row in rows:
        rec = dict(zip(header, row))
        if isinstance(rec.get("Plant ID"), int) and rec.get("Plant State") == "FL":
            generators.append(rec)
    wb.close()
    return generators


# ---------- 3. One row per plant ----------

def group_by_plant(generators: list[dict]) -> list[dict]:
    """Sum nameplate MW per plant; main technology = the one with the most MW there."""
    by_plant: dict[int, dict] = {}
    tech_mw: dict[int, Counter] = defaultdict(Counter)
    for g in generators:
        pid = g["Plant ID"]
        mw = g["Nameplate Capacity (MW)"]
        mw = float(mw) if isinstance(mw, (int, float)) else 0.0
        tech_mw[pid][g["Technology"]] += mw
        if pid not in by_plant:
            county = str(g["County"]).strip()
            by_plant[pid] = {
                "plant_id": pid,
                "name": str(g["Plant Name"]).strip(),
                "utility": str(g["Entity Name"]).strip(),
                "county": COUNTY_FIXES.get(county, county),
                "lat": round(float(g["Latitude"]), 5),
                "lon": round(float(g["Longitude"]), 5),
                "capacity_mw": 0.0,
            }
        by_plant[pid]["capacity_mw"] += mw
    for pid, plant in by_plant.items():
        plant["technology"] = tech_mw[pid].most_common(1)[0][0]
        plant["capacity_mw"] = round(plant["capacity_mw"], 1)
    return list(by_plant.values())


# ---------- 4. Pick 55 with statewide coverage ----------

def region(plant: dict) -> str:
    if plant["county"] in PANHANDLE_COUNTIES:
        return "panhandle"
    (lat_min, lat_max), (lon_min, lon_max) = EVERGLADES_BOX["lat"], EVERGLADES_BOX["lon"]
    if lat_min <= plant["lat"] < lat_max and lon_min < plant["lon"] < lon_max:
        return "everglades"
    if plant["lat"] >= 29.0:
        return "north"
    if plant["lat"] >= 27.5:
        return "central"
    return "south"


def miles_between(a: dict, b: dict) -> float:
    lat1, lat2 = math.radians(a["lat"]), math.radians(b["lat"])
    dlat, dlon = lat2 - lat1, math.radians(b["lon"] - a["lon"])
    h = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return 3958.8 * 2 * math.asin(math.sqrt(h))


def select(plants: list[dict], count: int) -> tuple[list[dict], dict[int, tuple[dict, dict]]]:
    """Largest plants by capacity, after first reserving the largest few in every region.

    Ties (many FPL solar plants are exactly 74.5 MW) break on plant_id, so reruns are stable.
    Returns (selected, skipped) where skipped maps plant_id -> (plant, the site it duplicates).
    """
    ranked = sorted(plants, key=lambda p: (-p["capacity_mw"], p["plant_id"]))
    chosen: dict[int, dict] = {}
    skipped: dict[int, tuple[dict, dict]] = {}

    def try_add(p: dict) -> bool:
        if p["plant_id"] in chosen or p["plant_id"] in skipped:
            return False
        near = next((c for c in chosen.values() if miles_between(c, p) < CO_LOCATED_MI), None)
        if near:
            skipped[p["plant_id"]] = (p, near)
            return False
        chosen[p["plant_id"]] = p
        return True

    for reg, minimum in REGION_MINIMUMS.items():
        added = 0
        for p in ranked:
            if added >= minimum:
                break
            if region(p) == reg and try_add(p):
                added += 1
    for p in ranked:
        if len(chosen) >= count:
            break
        try_add(p)
    selected = sorted(chosen.values(), key=lambda p: (-p["capacity_mw"], p["plant_id"]))
    return selected, skipped


# ---------- 5. Timezone and id ----------

def in_polygon(lon: float, lat: float, poly: list[tuple[float, float]]) -> bool:
    inside = False
    for (x1, y1), (x2, y2) in zip(poly, poly[1:] + poly[:1]):
        if (y1 > lat) != (y2 > lat) and lon < x1 + (lat - y1) * (x2 - x1) / (y2 - y1):
            inside = not inside
    return inside


def timezone_for(plant: dict) -> str:
    central = in_polygon(plant["lon"], plant["lat"], CENTRAL_TIME_POLYGON)
    if plant["county"] != "Gulf" and central != (plant["county"] in CENTRAL_COUNTIES):
        print(f"  WARNING timezone: {plant['name']} ({plant['county']}) polygon says "
              f"{'Central' if central else 'Eastern'}, county list disagrees")
    return "America/Chicago" if central else "America/New_York"


def slugify(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def assign_ids(sites: list[dict]) -> None:
    """id = slug of the plant name. Ids already in sites.json are kept (keyed by plant_id),
    so a plant renamed by EIA (Crist -> Gulf Clean Energy Center) keeps its id."""
    previous = {}
    if OUT_FILE.exists():
        previous = {s["plant_id"]: s["id"] for s in json.loads(OUT_FILE.read_text())["sites"]}
    taken = set()
    for s in sites:
        sid = previous.get(s["plant_id"]) or slugify(s["name"])
        if sid in taken:
            sid = f"{sid}-{s['plant_id']}"
        taken.add(sid)
        s["id"] = sid


# ---------- main ----------

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--file", type=Path, help="Use a local EIA-860M .xlsx instead of downloading")
    parser.add_argument("--count", type=int, default=TARGET_COUNT)
    args = parser.parse_args()

    workbook = args.file or download_latest()
    generators = read_fl_generators(workbook)
    plants = group_by_plant(generators)
    selected, skipped = select(plants, args.count)
    for s in selected:
        s["timezone"] = timezone_for(s)
    assign_ids(selected)

    keys = ["id", "plant_id", "name", "utility", "county", "lat", "lon", "timezone",
            "technology", "capacity_mw"]
    out = {
        "source": f"EIA-860M {workbook.name}, Operating sheet",
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "sites": [{k: s[k] for k in keys} for s in selected],
    }
    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    OUT_FILE.write_text(json.dumps(out, indent=2) + "\n")
    thermal = sorted((p for p in plants if p["technology"] not in NOT_THERMAL), key=lambda p: p["plant_id"])
    THERMAL_FILE.write_text(json.dumps({
        "source": out["source"],
        "plants": [{k: p[k] for k in ("plant_id", "name", "technology", "lat", "lon")} for p in thermal],
    }, indent=1) + "\n")

    central = sum(s["timezone"] == "America/Chicago" for s in selected)
    regions = Counter(region(s) for s in selected)
    print(f"\nFlorida plants found:  {len(plants)} ({len(generators)} operating generators)")
    print(f"Selected:              {len(selected)} ({sum(s['capacity_mw'] for s in selected):,.0f} MW)")
    print(f"On Central time:       {central}")
    print("By region:             " + ", ".join(f"{r} {regions[r]}" for r in REGION_MINIMUMS))
    for p, near in skipped.values():
        print(f"  skipped co-located {p['name']} ({miles_between(p, near):.2f} mi from {near['name']})")
    print(f"Thermal plants:        {len(thermal)} (for FIRMS plant hot spots)")
    print(f"Wrote {OUT_FILE.relative_to(BACKEND_DIR)} and {THERMAL_FILE.relative_to(BACKEND_DIR)}")


if __name__ == "__main__":
    main()
